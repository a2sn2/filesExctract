from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.comments import Comment
from pptx import Presentation
from reportlab.pdfgen import canvas


def run(*args: str) -> subprocess.CompletedProcess[str]:
    print(">", " ".join(args), flush=True)
    return subprocess.run(args, check=True, text=True, capture_output=True)


def main() -> int:
    if len(sys.argv) != 2:
        raise SystemExit("usage: smoke_windows.py <files-extract.exe>")

    exe = str(Path(sys.argv[1]).resolve())
    version = run(exe, "--version")
    assert "0.3.0" in version.stdout, version.stdout

    doctor = run(exe, "doctor", "--json")
    doctor_json = json.loads(doctor.stdout)
    assert doctor_json["dependencies"]["openpyxl"]["available"] is True
    assert doctor_json["dependencies"]["python-docx"]["available"] is True
    assert doctor_json["dependencies"]["python-pptx"]["available"] is True
    assert doctor_json["dependencies"]["pypdf"]["available"] is True

    with tempfile.TemporaryDirectory(prefix="files-extract-win-") as tmpdir:
        root = Path(tmpdir)

        xlsx = root / "sample.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "name"
        ws["B1"] = "amount"
        ws["A2"] = "alpha"
        ws["B2"] = 42
        ws["C2"] = "=B2*2"
        ws["A2"].comment = Comment("smoke-comment", "QA")
        ws["A2"].hyperlink = "https://example.com"
        wb.save(xlsx)

        docx = root / "sample.docx"
        doc = Document()
        doc.add_heading("Windows EXE smoke", level=1)
        doc.add_paragraph("Arabic: اختبار استخراج الملف")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "2"
        doc.save(docx)

        pptx = root / "sample.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Windows EXE smoke"
        slide.placeholders[1].text = "PowerPoint extraction works"
        prs.save(pptx)

        pdf = root / "sample.pdf"
        c = canvas.Canvas(str(pdf))
        c.drawString(72, 760, "Windows EXE PDF smoke")
        c.drawString(72, 740, "Text PDF fallback extraction")
        c.save()

        expected = {
            xlsx: "excel",
            docx: "word",
            pptx: "powerpoint",
            pdf: "pdf",
        }

        for source, expected_type in expected.items():
            inspected = run(exe, "inspect", str(source))
            assert f"type={expected_type}" in inspected.stdout, inspected.stdout
            out = root / f"{source.stem}-out"
            cmd = [exe, "extract", str(source), "-o", str(out)]
            if source.suffix.lower() == ".docx":
                cmd.append("--no-render-word-pages")
            extracted = run(*cmd)
            assert "Extracted:" in extracted.stdout, extracted.stdout
            payload = json.loads((out / "document.json").read_text(encoding="utf-8"))
            assert payload["metadata"]["document_type"] == expected_type, payload["metadata"]
            assert (out / "manifest.json").is_file()

    print("Windows executable smoke test: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

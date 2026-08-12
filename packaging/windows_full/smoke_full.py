from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas


def runtime_env(app: Path) -> dict[str, str]:
    env = os.environ.copy()
    lo = app / "tools" / "libreoffice" / "program"
    tess = app / "tools" / "tesseract"
    env["PATH"] = os.pathsep.join([str(lo), str(tess), env.get("PATH", "")])
    env["FILES_EXTRACT_BUNDLE_ROOT"] = str(app)
    env["TESSDATA_PREFIX"] = str(tess / "tessdata")
    env["DOCLING_ARTIFACTS_PATH"] = str(app / "models")
    env["DOCLING_SERVE_ARTIFACTS_PATH"] = str(app / "models")
    env["HF_HUB_OFFLINE"] = "1"
    env["TRANSFORMERS_OFFLINE"] = "1"
    env["HF_HUB_DISABLE_TELEMETRY"] = "1"
    env["PYTHONUTF8"] = "1"
    return env


def run_backend(app: Path, *args: str, check: bool = True) -> subprocess.CompletedProcess[str]:
    python = app / "runtime" / "python" / "python.exe"
    cmd = [str(python), "-m", "files_extract", *args]
    print(">", " ".join(cmd), flush=True)
    proc = subprocess.run(cmd, env=runtime_env(app), capture_output=True, text=True, check=False, timeout=420)
    if proc.stdout:
        print(proc.stdout)
    if proc.stderr:
        print(proc.stderr, file=sys.stderr)
    if check and proc.returncode != 0:
        raise AssertionError(f"command failed ({proc.returncode}): {cmd}\n{proc.stdout}\n{proc.stderr}")
    return proc


def create_fixtures(root: Path) -> dict[str, Path]:
    xlsx = root / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "name"
    ws["B1"] = "amount"
    ws["A2"] = "alpha"
    ws["B2"] = 42
    ws["C2"] = "=B2*2"
    wb.create_sheet("Hidden").sheet_state = "hidden"
    wb.save(xlsx)

    docx = root / "sample.docx"
    doc = Document()
    doc.add_heading("Windows full smoke", level=1)
    doc.add_paragraph("Arabic text: اختبار استخراج المستند")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.add_page_break()
    doc.add_paragraph("Second rendered page")
    doc.save(docx)

    pptx = root / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Windows full smoke"
    slide.placeholders[1].text = "PowerPoint extraction works"
    prs.save(pptx)

    pdf = root / "sample.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(72, 760, "Windows full PDF smoke")
    c.drawString(72, 740, "Text and layout extraction")
    c.save()

    scanned = root / "scanned.pdf"
    image = Image.new("RGB", (1800, 1000), "white")
    draw = ImageDraw.Draw(image)
    font = None
    for candidate in [Path(r"C:\Windows\Fonts\arial.ttf"), Path(r"C:\Windows\Fonts\segoeui.ttf")]:
        if candidate.exists():
            font = ImageFont.truetype(str(candidate), 86)
            break
    if font is None:
        font = ImageFont.load_default()
    draw.text((120, 360), "OFFLINE OCR TEST 12345", fill="black", font=font)
    image.save(scanned, "PDF", resolution=150.0)

    encrypted = root / "encrypted.pdf"
    reader = PdfReader(pdf)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("secret-test-password")
    with encrypted.open("wb") as handle:
        writer.write(handle)

    return {"xlsx": xlsx, "docx": docx, "pptx": pptx, "pdf": pdf, "scanned": scanned, "encrypted": encrypted}


def json_payload(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def assert_primary_docling(payload: dict) -> None:
    pdf = payload.get("metadata", {}).get("properties", {}).get("pdf", {})
    assert pdf.get("primary_backend") == "docling", pdf
    ocr = pdf.get("ocr", {})
    assert ocr.get("enabled") is True, ocr
    layout = pdf.get("layout_engine", {})
    assert layout.get("runtime") == "onnxruntime", layout
    assert layout.get("requires_runtime_compiler") is False, layout


def libreoffice_convert(app: Path, source: Path, extension: str, out: Path) -> Path:
    soffice = app / "tools" / "libreoffice" / "program" / "soffice.exe"
    out.mkdir(exist_ok=True)
    cmd = [str(soffice), "--headless", "--convert-to", extension, "--outdir", str(out), str(source)]
    proc = subprocess.run(cmd, env=runtime_env(app), capture_output=True, text=True, timeout=180, check=False)
    if proc.returncode != 0:
        raise AssertionError(f"LibreOffice conversion failed: {proc.stdout}\n{proc.stderr}")
    expected = out / f"{source.stem}.{extension}"
    if not expected.exists():
        candidates = list(out.glob(f"*.{extension}"))
        assert candidates, f"No .{extension} output from LibreOffice. {proc.stdout} {proc.stderr}"
        return candidates[0]
    return expected


def main() -> int:
    if len(sys.argv) != 2:
        raise SystemExit("usage: smoke_full.py <app-root>")
    app = Path(sys.argv[1]).resolve()
    assert (app / "FilesExtract.exe").is_file(), app
    assert (app / "runtime" / "python" / "python.exe").is_file()

    version = run_backend(app, "--version")
    assert "0.4.0" in version.stdout

    doctor = json.loads(run_backend(app, "doctor", "--json").stdout)
    assert doctor["dependencies"]["openpyxl"]["available"] is True
    assert doctor["dependencies"]["python-docx"]["available"] is True
    assert doctor["dependencies"]["python-pptx"]["available"] is True
    assert doctor["dependencies"]["pypdf"]["available"] is True
    assert doctor["dependencies"]["docling"]["available"] is True
    lo_path = Path(doctor["external_tools"]["libreoffice"]["path"]).resolve()
    tess_path = Path(doctor["external_tools"]["tesseract"]["path"]).resolve()
    assert app in lo_path.parents, lo_path
    assert app in tess_path.parents, tess_path
    languages = set(doctor["external_tools"]["tesseract"].get("languages") or [])
    assert {"ara", "eng"} <= languages, languages

    with tempfile.TemporaryDirectory(prefix="files-extract-full-smoke-") as tmp:
        root = Path(tmp)
        fixtures = create_fixtures(root)
        output = root / "output"
        output.mkdir()

        expected_types = {"xlsx": "xlsx", "docx": "docx", "pptx": "pptx", "pdf": "pdf"}
        for key, expected in expected_types.items():
            destination = output / key
            run_backend(app, "extract", str(fixtures[key]), "-o", str(destination))
            payload = json_payload(destination / "document.json")
            assert payload["metadata"]["document_type"] == expected
            assert (destination / "manifest.json").is_file()

        word_payload = json_payload(output / "docx" / "document.json")
        pagination = word_payload.get("metadata", {}).get("properties", {}).get("word", {}).get("pagination", {})
        assert pagination, "Word rendered pagination metadata missing"
        assert int(pagination.get("rendered_page_count") or 0) >= 2, pagination

        assert_primary_docling(json_payload(output / "pdf" / "document.json"))

        scanned_out = output / "scanned"
        run_backend(app, "extract", str(fixtures["scanned"]), "-o", str(scanned_out))
        scanned_payload = json_payload(scanned_out / "document.json")
        assert_primary_docling(scanned_payload)
        all_text = "\n".join(
            (element.get("text") or "")
            for unit in scanned_payload.get("units", [])
            for element in unit.get("elements", [])
        ).upper()
        assert "OCR" in all_text and "12345" in all_text, all_text[:1000]

        encrypted_out = output / "encrypted"
        run_backend(app, "extract", str(fixtures["encrypted"]), "-o", str(encrypted_out), "--password", "secret-test-password")
        assert (encrypted_out / "document.json").is_file()

        legacy_dir = root / "legacy"
        legacy = {
            "doc": libreoffice_convert(app, fixtures["docx"], "doc", legacy_dir),
            "xls": libreoffice_convert(app, fixtures["xlsx"], "xls", legacy_dir),
            "ppt": libreoffice_convert(app, fixtures["pptx"], "ppt", legacy_dir),
        }
        for extension, source in legacy.items():
            dest = output / f"legacy-{extension}"
            run_backend(app, "extract", str(source), "-o", str(dest))
            assert (dest / "document.json").is_file()

        launcher = subprocess.run(
            [str(app / "FilesExtract.exe"), "--version"],
            env=runtime_env(app),
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        assert launcher.returncode == 0, launcher.stderr

    print("WINDOWS FULL PRODUCT SMOKE: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

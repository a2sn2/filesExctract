from __future__ import annotations

from pathlib import Path

import pytest
from msoffcrypto.format.ooxml import OOXMLFile
from openpyxl import Workbook

from files_extract.detection import detect_file_type
from files_extract.engine import ExtractionEngine, ExtractionOptions
from files_extract.errors import ExtractionFailedError


def encrypted_xlsx(tmp_path: Path, password: str = "office-secret") -> Path:
    plain = tmp_path / "plain.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Protected Data"
    ws["A1"] = "account"
    ws["B1"] = "amount"
    ws["A2"] = "alpha"
    ws["B2"] = 123.45
    wb.save(plain)

    encrypted = tmp_path / "protected.xlsx"
    with plain.open("rb") as source, encrypted.open("wb") as destination:
        OOXMLFile(source).encrypt(password, destination)
    return encrypted


def test_encrypted_modern_office_is_decrypted_and_extracted(tmp_path: Path) -> None:
    source = encrypted_xlsx(tmp_path)
    detected = detect_file_type(source)
    assert detected.is_encrypted_office is True
    assert detected.extension == "xlsx"

    engine = ExtractionEngine(
        options=ExtractionOptions(
            render_word_pages=False,
            pdf_password="office-secret",
        )
    )
    document = engine.extract(source)

    assert document.metadata.filename == "protected.xlsx"
    assert document.metadata.document_type == "xlsx"
    assert document.metadata.properties["source"]["is_encrypted_office"] is True
    security = document.metadata.properties["security"]
    assert security["encrypted_source"] is True
    assert security["decrypted_for_extraction"] is True
    assert security["password_persisted"] is False
    assert any(unit.name == "Protected Data" for unit in document.units)
    assert any(w.code == "office.encrypted.decrypted" for w in document.warnings)


def test_encrypted_modern_office_requires_password(tmp_path: Path) -> None:
    source = encrypted_xlsx(tmp_path)
    engine = ExtractionEngine(options=ExtractionOptions(render_word_pages=False))

    with pytest.raises(ExtractionFailedError, match="no password"):
        engine.extract(source)


def test_encrypted_modern_office_rejects_wrong_password(tmp_path: Path) -> None:
    source = encrypted_xlsx(tmp_path)
    engine = ExtractionEngine(
        options=ExtractionOptions(
            render_word_pages=False,
            pdf_password="wrong-password",
        )
    )

    with pytest.raises(ExtractionFailedError, match="Unable to decrypt"):
        engine.extract(source)

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from files_extract.engine import ExtractionEngine
from files_extract.models import ElementType, UnitKind
from files_extract.renderers import render_json, render_markdown


def _build_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Product", "Amount", "URL", "Merged"])
    sheet.append(["A", 10, "OpenAI", "Header"])
    sheet.append(["B", 20, None, None])
    sheet.append(["Total", "=SUM(B2:B3)", None, None])
    sheet["C2"].hyperlink = "https://example.com"
    sheet["C2"].comment = Comment("sample comment", "tester")
    sheet.merge_cells("D2:E2")
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions["C"].hidden = True
    sheet.freeze_panes = "A2"

    table = Table(displayName="SalesTable", ref="A1:B4")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)

    validation = DataValidation(type="whole", operator="between", formula1="1", formula2="100")
    sheet.add_data_validation(validation)
    validation.add("B2:B3")

    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "secret"

    workbook.properties.creator = "files-extract test"
    workbook.save(path)
    workbook.close()


def test_extracts_xlsx_structure_and_cell_metadata(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _build_workbook(path)

    document = ExtractionEngine().extract(path)

    assert document.metadata.document_type == "xlsx"
    assert document.unit_count == 2
    assert document.units[0].kind == UnitKind.SHEET
    assert document.units[0].name == "Sales"
    assert document.units[1].metadata["state"] == "hidden"

    metadata = document.units[0].metadata
    assert "D2:E2" in metadata["merged_ranges"]
    assert 3 in metadata["hidden_rows"]
    assert "C" in metadata["hidden_columns"]
    assert metadata["tables"][0]["display_name"] == "SalesTable"
    assert metadata["data_validations"][0]["ranges"] == "B2:B3"

    cells = {
        element.data["coordinate"]: element
        for element in document.units[0].elements
        if element.element_type == ElementType.CELL
    }
    assert cells["B4"].data["formula"] == "=SUM(B2:B3)"
    assert cells["C2"].data["hyperlink"]["target"] == "https://example.com"
    assert cells["C2"].data["comment"]["text"] == "sample comment"

    json_output = render_json(document)
    markdown_output = render_markdown(document)
    assert '"formula": "=SUM(B2:B3)"' in json_output
    assert "## Sheet 1: Sales" in markdown_output
    assert "sample comment" in markdown_output

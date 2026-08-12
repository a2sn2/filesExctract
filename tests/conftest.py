from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image
from docx import Document
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from reportlab.pdfgen import canvas


@pytest.fixture
def sample_image(tmp_path: Path) -> Path:
    path = tmp_path / "sample-image.png"
    Image.new("RGB", (64, 32), (255, 255, 255)).save(path)
    return path


@pytest.fixture
def xlsx_file(tmp_path: Path, sample_image: Path) -> Path:
    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "بيانات"
    ws.append(["Name", "Qty", "Total"])
    ws.append(["A", 2, "=B2*10"])
    ws["A2"].comment = Comment("ملاحظة", "QA")
    ws["A2"].hyperlink = "https://example.com"
    ws.merge_cells("A4:B4")
    ws["A4"] = "merged"
    ws.row_dimensions[5].hidden = True
    ws.column_dimensions["D"].hidden = True
    tab = Table(displayName="SalesTable", ref="A1:C2")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    from openpyxl.drawing.image import Image as XLImage
    ws.add_image(XLImage(sample_image), "E2")
    wb.create_sheet("Hidden").sheet_state = "hidden"
    wb.save(path)
    return path


@pytest.fixture
def docx_file(tmp_path: Path, sample_image: Path) -> Path:
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.core_properties.title = "اختبار Word"
    doc.add_heading("العنوان الرئيسي", level=1)
    p = doc.add_paragraph("فقرة عربية للاختبار")
    p.runs[0].bold = True
    doc.add_comment(p.runs, text="تعليق مراجعة", author="QA", initials="QA")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "nested"
    doc.add_picture(str(sample_image), width=Inches(1))
    doc.sections[0].header.paragraphs[0].text = "رأس الصفحة"
    doc.sections[0].footer.paragraphs[0].text = "تذييل الصفحة"
    doc.add_page_break()
    doc.add_paragraph("الصفحة الثانية")
    doc.save(path)
    return path


@pytest.fixture
def pptx_file(tmp_path: Path, sample_image: Path) -> Path:
    path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "عنوان العرض"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1.2), Inches(4), Inches(0.8))
    textbox.text_frame.text = "نص داخل الشريحة"
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1.2))
    table_shape.table.cell(0, 0).text = "A"
    table_shape.table.cell(0, 1).text = "B"
    table_shape.table.cell(1, 0).text = "1"
    table_shape.table.cell(1, 1).text = "2"
    slide.shapes.add_picture(str(sample_image), Inches(5), Inches(1), width=Inches(1.2))
    data = ChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Sales", (10, 20))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3.5), Inches(4), Inches(2), data)
    slide.notes_slide.notes_text_frame.text = "ملاحظات المتحدث"
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Hidden slide"
    slide2._element.set("show", "0")
    prs.save(path)
    return path


@pytest.fixture
def pdf_file(tmp_path: Path, sample_image: Path) -> Path:
    path = tmp_path / "sample.pdf"
    c = canvas.Canvas(str(path))
    c.setTitle("PDF Test")
    c.drawString(72, 750, "Hello PDF page one")
    c.linkURL("https://example.com", (72, 730, 200, 745), relative=0)
    c.drawImage(str(sample_image), 72, 650, width=100, height=50)
    c.showPage()
    c.drawString(72, 750, "Second PDF page")
    c.save()
    return path

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from files_extract.engine import ExtractionEngine
from files_extract.libreoffice import convert_office, find_libreoffice


@pytest.mark.integration
def test_word_rendered_page_count(docx_file: Path) -> None:
    if not find_libreoffice():
        pytest.skip("LibreOffice not installed")
    doc = ExtractionEngine().extract(docx_file)
    assert doc.metadata.properties["word"]["pagination"]["rendered_page_count"] >= 2


@pytest.mark.integration
def test_legacy_doc_conversion(tmp_path: Path) -> None:
    if not find_libreoffice():
        pytest.skip("LibreOffice not installed")
    modern = tmp_path / "legacy-source.docx"
    d = Document(); d.add_paragraph("Legacy document content"); d.save(modern)
    legacy = convert_office(modern, "doc", tmp_path / "legacy")
    extracted = ExtractionEngine().extract(legacy)
    assert extracted.metadata.document_type == "doc"
    assert any(e.text == "Legacy document content" for e in extracted.units[0].elements)

@pytest.mark.integration
def test_legacy_doc_embedded_asset_materialization(tmp_path: Path, sample_image: Path) -> None:
    if not find_libreoffice():
        pytest.skip("LibreOffice not installed")

    modern = tmp_path / "legacy-with-image.docx"
    document = Document()
    document.add_paragraph("Legacy asset")
    document.add_picture(str(sample_image))
    document.save(modern)

    legacy = convert_office(modern, "doc", tmp_path / "legacy-assets")
    extracted = ExtractionEngine().extract(legacy)
    assert extracted.assets
    assert all(a.metadata.get("legacy_conversion_target") == "docx" for a in extracted.assets)

    from files_extract.output import write_output_package

    output_dir = tmp_path / "legacy-output"
    result = write_output_package(legacy, extracted, output_dir)
    assert result["summary"]["assets_written"] > 0
    asset_entries = [item for item in result["files"] if str(item["path"]).startswith("assets/")]
    assert asset_entries
    assert all((output_dir / str(item["path"])).exists() for item in asset_entries)


@pytest.mark.integration
def test_legacy_xls_conversion(tmp_path: Path, xlsx_file: Path) -> None:
    if not find_libreoffice():
        pytest.skip("LibreOffice not installed")
    legacy = convert_office(xlsx_file, "xls", tmp_path / "legacy-xls")
    extracted = ExtractionEngine().extract(legacy)
    assert extracted.metadata.document_type == "xls"
    assert extracted.units[0].name == "بيانات"
    assert any(e.data.get("coordinate") == "A1" for e in extracted.units[0].elements)


@pytest.mark.integration
def test_legacy_ppt_conversion(tmp_path: Path, pptx_file: Path) -> None:
    if not find_libreoffice():
        pytest.skip("LibreOffice not installed")
    legacy = convert_office(pptx_file, "ppt", tmp_path / "legacy-ppt")
    extracted = ExtractionEngine().extract(legacy)
    assert extracted.metadata.document_type == "ppt"
    text = "\n".join(e.text or "" for u in extracted.units for e in u.elements)
    assert "عنوان العرض" in text

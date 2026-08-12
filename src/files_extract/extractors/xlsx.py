from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from ..detection import DetectedFileType
from ..errors import ExtractionFailedError
from ..models import (
    CanonicalDocument,
    DocumentElement,
    DocumentMetadata,
    DocumentUnit,
    ElementType,
    ExtractionWarning,
    Severity,
    SourceReference,
    UnitKind,
    UnsupportedObject,
)
from .base import BaseExtractor


def _safe_attr(obj: Any, name: str, default: Any = None) -> Any:
    try:
        return getattr(obj, name, default)
    except Exception:
        return default


def _anchor_metadata(anchor: Any) -> dict[str, Any]:
    if isinstance(anchor, str):
        return {"anchor": anchor}
    result: dict[str, Any] = {}
    start = _safe_attr(anchor, "_from")
    end = _safe_attr(anchor, "to")
    if start is not None:
        result["from"] = {
            "column_zero_based": _safe_attr(start, "col"),
            "row_zero_based": _safe_attr(start, "row"),
            "column_offset": _safe_attr(start, "colOff"),
            "row_offset": _safe_attr(start, "rowOff"),
        }
    if end is not None:
        result["to"] = {
            "column_zero_based": _safe_attr(end, "col"),
            "row_zero_based": _safe_attr(end, "row"),
            "column_offset": _safe_attr(end, "colOff"),
            "row_offset": _safe_attr(end, "rowOff"),
        }
    return result


class XlsxExtractor(BaseExtractor):
    name = "openpyxl-native"
    supported_types = frozenset({"xlsx", "xlsm"})

    def __init__(self) -> None:
        try:
            import openpyxl

            self.version = openpyxl.__version__
        except Exception:
            self.version = None

    def extract(
        self,
        path: Path,
        detected: DetectedFileType,
        metadata: DocumentMetadata,
    ) -> CanonicalDocument:
        keep_vba = detected.document_type == "xlsm"
        try:
            workbook = load_workbook(
                path,
                data_only=False,
                read_only=False,
                keep_vba=keep_vba,
                keep_links=True,
            )
            cached_workbook = load_workbook(
                path,
                data_only=True,
                read_only=False,
                keep_vba=keep_vba,
                keep_links=True,
            )
        except Exception as exc:
            raise ExtractionFailedError(f"Failed to open workbook: {path.name}: {exc}") from exc

        warnings: list[ExtractionWarning] = []
        unsupported: list[UnsupportedObject] = []
        units: list[DocumentUnit] = []

        try:
            metadata.properties.update(self._workbook_metadata(workbook))

            if keep_vba and _safe_attr(workbook, "vba_archive") is not None:
                unsupported.append(
                    UnsupportedObject(
                        object_type="vba_project",
                        status="detected_not_parsed",
                        description="The VBA project is preserved by openpyxl but its source code is not parsed yet.",
                        metadata={"container": "xlsm"},
                    )
                )

            external_links = list(_safe_attr(workbook, "_external_links", []) or [])
            if external_links:
                unsupported.append(
                    UnsupportedObject(
                        object_type="external_workbook_links",
                        status="detected_partially_parsed",
                        description="External workbook link objects were detected; detailed external content is not resolved.",
                        metadata={"count": len(external_links)},
                    )
                )

            for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
                cached_sheet = cached_workbook[worksheet.title]
                unit, sheet_warnings, sheet_unsupported = self._extract_sheet(
                    worksheet,
                    cached_sheet,
                    sheet_index,
                )
                units.append(unit)
                warnings.extend(sheet_warnings)
                unsupported.extend(sheet_unsupported)
        finally:
            workbook.close()
            cached_workbook.close()

        if not units:
            warnings.append(
                ExtractionWarning(
                    code="xlsx.no_sheets",
                    message="Workbook contains no worksheets.",
                    severity=Severity.WARNING,
                )
            )

        return CanonicalDocument(
            metadata=metadata,
            units=units,
            warnings=warnings,
            unsupported_objects=unsupported,
        )

    def _workbook_metadata(self, workbook: Any) -> dict[str, Any]:
        properties = workbook.properties
        property_names = (
            "title",
            "subject",
            "creator",
            "keywords",
            "description",
            "lastModifiedBy",
            "created",
            "modified",
            "category",
            "contentStatus",
            "identifier",
            "language",
            "version",
            "revision",
            "lastPrinted",
        )
        core = {
            name: _safe_attr(properties, name)
            for name in property_names
            if _safe_attr(properties, name) is not None
        }

        defined_names: list[dict[str, Any]] = []
        try:
            for item in workbook.defined_names.values():
                defined_names.append(
                    {
                        "name": _safe_attr(item, "name"),
                        "value": _safe_attr(item, "attr_text"),
                        "type": _safe_attr(item, "type"),
                        "hidden": _safe_attr(item, "hidden"),
                        "comment": _safe_attr(item, "comment"),
                        "local_sheet_id": _safe_attr(item, "localSheetId"),
                    }
                )
        except Exception:
            defined_names = []

        calculation = _safe_attr(workbook, "calculation")
        calculation_metadata = {}
        if calculation is not None:
            for name in ("calcMode", "calcId", "fullCalcOnLoad", "forceFullCalc", "iterate"):
                value = _safe_attr(calculation, name)
                if value is not None:
                    calculation_metadata[name] = value

        return {
            "workbook": {
                "core_properties": core,
                "sheet_names": list(workbook.sheetnames),
                "defined_names": defined_names,
                "calculation": calculation_metadata,
            }
        }

    def _extract_sheet(
        self,
        worksheet: Any,
        cached_sheet: Any,
        sheet_index: int,
    ) -> tuple[DocumentUnit, list[ExtractionWarning], list[UnsupportedObject]]:
        source_base = SourceReference(
            unit_index=sheet_index,
            unit_name=worksheet.title,
            sheet_name=worksheet.title,
        )
        warnings: list[ExtractionWarning] = []
        unsupported: list[UnsupportedObject] = []

        merged_ranges = [str(item) for item in worksheet.merged_cells.ranges]
        hidden_rows = sorted(
            index
            for index, dimension in worksheet.row_dimensions.items()
            if _safe_attr(dimension, "hidden") is True
        )
        hidden_columns = sorted(
            key
            for key, dimension in worksheet.column_dimensions.items()
            if _safe_attr(dimension, "hidden") is True
        )

        tables: list[dict[str, Any]] = []
        try:
            for table in worksheet.tables.values():
                style_info = _safe_attr(table, "tableStyleInfo")
                tables.append(
                    {
                        "name": _safe_attr(table, "name"),
                        "display_name": _safe_attr(table, "displayName"),
                        "ref": _safe_attr(table, "ref"),
                        "header_row_count": _safe_attr(table, "headerRowCount"),
                        "totals_row_count": _safe_attr(table, "totalsRowCount"),
                        "style": {
                            "name": _safe_attr(style_info, "name"),
                            "show_first_column": _safe_attr(style_info, "showFirstColumn"),
                            "show_last_column": _safe_attr(style_info, "showLastColumn"),
                            "show_row_stripes": _safe_attr(style_info, "showRowStripes"),
                            "show_column_stripes": _safe_attr(style_info, "showColumnStripes"),
                        }
                        if style_info is not None
                        else None,
                    }
                )
        except Exception as exc:
            warnings.append(
                ExtractionWarning(
                    code="xlsx.tables.partial",
                    message="One or more Excel table definitions could not be fully read.",
                    source=source_base,
                    details={"error": str(exc)},
                )
            )

        data_validations: list[dict[str, Any]] = []
        try:
            validations = _safe_attr(_safe_attr(worksheet, "data_validations"), "dataValidation", []) or []
            for validation in validations:
                data_validations.append(
                    {
                        "type": _safe_attr(validation, "type"),
                        "operator": _safe_attr(validation, "operator"),
                        "ranges": str(_safe_attr(validation, "sqref", "")),
                        "formula1": _safe_attr(validation, "formula1"),
                        "formula2": _safe_attr(validation, "formula2"),
                        "allow_blank": _safe_attr(validation, "allowBlank"),
                        "show_error_message": _safe_attr(validation, "showErrorMessage"),
                        "show_input_message": _safe_attr(validation, "showInputMessage"),
                        "error": _safe_attr(validation, "error"),
                        "error_title": _safe_attr(validation, "errorTitle"),
                        "prompt": _safe_attr(validation, "prompt"),
                        "prompt_title": _safe_attr(validation, "promptTitle"),
                    }
                )
        except Exception as exc:
            warnings.append(
                ExtractionWarning(
                    code="xlsx.data_validations.partial",
                    message="One or more data validation rules could not be fully read.",
                    source=source_base,
                    details={"error": str(exc)},
                )
            )

        sheet_metadata = {
            "state": worksheet.sheet_state,
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "merged_ranges": merged_ranges,
            "hidden_rows": hidden_rows,
            "hidden_columns": hidden_columns,
            "freeze_panes": str(worksheet.freeze_panes) if worksheet.freeze_panes else None,
            "auto_filter": _safe_attr(worksheet.auto_filter, "ref"),
            "print_area": str(worksheet.print_area) if worksheet.print_area else None,
            "print_title_rows": _safe_attr(worksheet, "print_title_rows"),
            "print_title_cols": _safe_attr(worksheet, "print_title_cols"),
            "tables": tables,
            "data_validations": data_validations,
        }

        elements: list[DocumentElement] = [
            DocumentElement(
                element_id=f"sheet-{sheet_index}-metadata",
                element_type=ElementType.METADATA,
                order=0,
                data=sheet_metadata,
                source=source_base,
            )
        ]

        order = 1
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if not self._should_emit_cell(cell):
                    continue

                cached_cell = cached_sheet[cell.coordinate]
                source = SourceReference(
                    unit_index=sheet_index,
                    unit_name=worksheet.title,
                    sheet_name=worksheet.title,
                    row=cell.row,
                    column=cell.column,
                    cell=cell.coordinate,
                )
                formula = cell.value if cell.data_type == "f" else None
                cached_value = cached_cell.value if formula is not None else None
                hyperlink = self._hyperlink_data(cell)
                comment = self._comment_data(cell)

                value = cell.value
                cell_data = {
                    "coordinate": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "value": value,
                    "formula": formula,
                    "cached_value": cached_value,
                    "data_type": cell.data_type,
                    "number_format": cell.number_format,
                    "style_id": cell.style_id,
                    "is_date": bool(_safe_attr(cell, "is_date", False)),
                    "hyperlink": hyperlink,
                    "comment": comment,
                    "row_hidden": cell.row in hidden_rows,
                    "column_hidden": cell.column_letter in hidden_columns,
                }
                elements.append(
                    DocumentElement(
                        element_id=f"sheet-{sheet_index}-cell-{cell.coordinate}",
                        element_type=ElementType.CELL,
                        order=order,
                        text=self._cell_text(value),
                        data=cell_data,
                        source=source,
                    )
                )
                order += 1

        for table_index, table_data in enumerate(tables, start=1):
            elements.append(
                DocumentElement(
                    element_id=f"sheet-{sheet_index}-table-{table_index}",
                    element_type=ElementType.TABLE,
                    order=order,
                    data=table_data,
                    source=source_base,
                )
            )
            order += 1

        images = list(_safe_attr(worksheet, "_images", []) or [])
        for image_index, image in enumerate(images, start=1):
            unsupported.append(
                UnsupportedObject(
                    object_type="xlsx_image",
                    status="detected_not_extracted",
                    description="An embedded worksheet image was detected. Asset byte extraction is scheduled for the asset phase.",
                    source=source_base,
                    metadata={
                        "index": image_index,
                        "format": _safe_attr(image, "format"),
                        "width": _safe_attr(image, "width"),
                        "height": _safe_attr(image, "height"),
                        **_anchor_metadata(_safe_attr(image, "anchor")),
                    },
                )
            )

        charts = list(_safe_attr(worksheet, "_charts", []) or [])
        for chart_index, chart in enumerate(charts, start=1):
            unsupported.append(
                UnsupportedObject(
                    object_type="xlsx_chart",
                    status="detected_partially_parsed",
                    description="A worksheet chart was detected. Chart rendering and full series extraction are not implemented yet.",
                    source=source_base,
                    metadata={
                        "index": chart_index,
                        "title": str(_safe_attr(chart, "title")) if _safe_attr(chart, "title") is not None else None,
                        "style": _safe_attr(chart, "style"),
                        **_anchor_metadata(_safe_attr(chart, "anchor")),
                    },
                )
            )

        unit = DocumentUnit(
            index=sheet_index,
            kind=UnitKind.SHEET,
            name=worksheet.title,
            elements=elements,
            metadata=sheet_metadata,
        )
        return unit, warnings, unsupported

    @staticmethod
    def _should_emit_cell(cell: Any) -> bool:
        return (
            cell.value is not None
            or _safe_attr(cell, "comment") is not None
            or _safe_attr(cell, "hyperlink") is not None
        )

    @staticmethod
    def _cell_text(value: Any) -> str | None:
        if value is None:
            return None
        return str(value)

    @staticmethod
    def _hyperlink_data(cell: Any) -> dict[str, Any] | None:
        link = _safe_attr(cell, "hyperlink")
        if link is None:
            return None
        return {
            "target": _safe_attr(link, "target"),
            "location": _safe_attr(link, "location"),
            "display": _safe_attr(link, "display"),
            "tooltip": _safe_attr(link, "tooltip"),
            "id": _safe_attr(link, "id"),
        }

    @staticmethod
    def _comment_data(cell: Any) -> dict[str, Any] | None:
        comment = _safe_attr(cell, "comment")
        if comment is None:
            return None
        return {
            "text": _safe_attr(comment, "text"),
            "author": _safe_attr(comment, "author"),
            "width": _safe_attr(comment, "width"),
            "height": _safe_attr(comment, "height"),
        }

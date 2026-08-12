from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation

from ..assets import collect_ooxml_assets
from ..detection import DetectedFileType
from ..errors import ExtractionFailedError
from ..models import (
    CanonicalDocument,
    DocumentElement,
    DocumentMetadata,
    DocumentUnit,
    ElementType,
    ExtractionWarning,
    SourceReference,
    UnitKind,
    UnsupportedObject,
)
from .base import BaseExtractor


def _safe_attr(obj: Any, name: str, default: Any = None) -> Any:
    try:
        return getattr(obj, name, default)
    except Exception:
        return default


def _enum_name(value: Any) -> str | None:
    if value is None:
        return None
    return getattr(value, "name", str(value))


def _length(value: Any) -> int | None:
    return int(value) if value is not None else None


def _core_properties(props: Any) -> dict[str, Any]:
    names = (
        "author", "category", "comments", "content_status", "created", "identifier", "keywords",
        "language", "last_modified_by", "last_printed", "modified", "revision", "subject", "title", "version",
    )
    return {name: _safe_attr(props, name) for name in names if _safe_attr(props, name) not in (None, "")}


def _run_data(run: Any) -> dict[str, Any]:
    font = run.font
    hyperlink = _safe_attr(run, "hyperlink")
    return {
        "text": run.text,
        "font": {
            "name": font.name,
            "size_emu": _length(font.size),
            "bold": font.bold,
            "italic": font.italic,
            "underline": str(font.underline) if font.underline is not None else None,
        },
        "hyperlink": _safe_attr(hyperlink, "address") if hyperlink is not None else None,
    }


def _paragraph_data(paragraph: Any) -> dict[str, Any]:
    font = paragraph.font
    return {
        "text": paragraph.text,
        "level": paragraph.level,
        "alignment": _enum_name(paragraph.alignment),
        "line_spacing": str(paragraph.line_spacing) if paragraph.line_spacing is not None else None,
        "space_before": _length(paragraph.space_before),
        "space_after": _length(paragraph.space_after),
        "font": {
            "name": font.name,
            "size_emu": _length(font.size),
            "bold": font.bold,
            "italic": font.italic,
        },
        "runs": [_run_data(run) for run in paragraph.runs],
    }


def _shape_geometry(shape: Any) -> dict[str, Any]:
    return {
        "left": _length(_safe_attr(shape, "left")),
        "top": _length(_safe_attr(shape, "top")),
        "width": _length(_safe_attr(shape, "width")),
        "height": _length(_safe_attr(shape, "height")),
        "rotation": _safe_attr(shape, "rotation"),
    }


def _shape_accessibility(shape: Any) -> dict[str, Any]:
    try:
        element = shape._element
        for node in element.iter():
            if str(node.tag).endswith("}cNvPr"):
                return {
                    "title": node.get("title"),
                    "description": node.get("descr"),
                    "name": node.get("name"),
                }
    except Exception:
        pass
    return {"title": None, "description": None, "name": None}


def _shape_text_snapshot(shapes: Iterable[Any]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for index, shape in enumerate(shapes, start=1):
        text, _ = _shape_text(shape)
        accessibility = _shape_accessibility(shape)
        if text or accessibility.get("description") or accessibility.get("title"):
            result.append({
                "index": index,
                "name": _safe_attr(shape, "name"),
                "text": text,
                "accessibility": accessibility,
            })
    return result


def _placeholder_data(shape: Any) -> dict[str, Any] | None:
    if not bool(_safe_attr(shape, "is_placeholder", False)):
        return None
    try:
        fmt = shape.placeholder_format
        return {"idx": fmt.idx, "type": _enum_name(fmt.type)}
    except Exception:
        return {"idx": None, "type": None}


def _table_data(table: Any) -> dict[str, Any]:
    rows: list[list[dict[str, Any]]] = []
    for row_idx, row in enumerate(table.rows):
        cells: list[dict[str, Any]] = []
        for col_idx, cell in enumerate(row.cells):
            cells.append({
                "row": row_idx,
                "column": col_idx,
                "text": cell.text,
                "is_merge_origin": bool(_safe_attr(cell, "is_merge_origin", False)),
                "is_spanned": bool(_safe_attr(cell, "is_spanned", False)),
                "span_width": _safe_attr(cell, "span_width"),
                "span_height": _safe_attr(cell, "span_height"),
                "margin_left": _length(_safe_attr(cell, "margin_left")),
                "margin_right": _length(_safe_attr(cell, "margin_right")),
                "margin_top": _length(_safe_attr(cell, "margin_top")),
                "margin_bottom": _length(_safe_attr(cell, "margin_bottom")),
                "vertical_anchor": _enum_name(_safe_attr(cell, "vertical_anchor")),
                "paragraphs": [_paragraph_data(p) for p in cell.text_frame.paragraphs],
            })
        rows.append(cells)
    return {
        "rows": rows,
        "row_count": len(table.rows),
        "column_count": len(table.columns),
        "first_row": table.first_row,
        "last_row": table.last_row,
        "first_col": table.first_col,
        "last_col": table.last_col,
        "horz_banding": table.horz_banding,
        "vert_banding": table.vert_banding,
    }


def _chart_data(chart: Any) -> dict[str, Any]:
    title = None
    if bool(_safe_attr(chart, "has_title", False)):
        try:
            title = chart.chart_title.text_frame.text
        except Exception:
            title = None
    series: list[dict[str, Any]] = []
    try:
        for item in chart.series:
            series.append({"name": item.name, "values": list(item.values)})
    except Exception:
        pass
    categories: list[str] = []
    try:
        if chart.plots:
            categories = [str(_safe_attr(cat, "label", "")) for cat in chart.plots[0].categories]
    except Exception:
        pass
    return {
        "title": title,
        "series": series,
        "categories": categories,
        "has_legend": bool(_safe_attr(chart, "has_legend", False)),
        "chart_style": _safe_attr(chart, "chart_style"),
    }


def _shape_text(shape: Any) -> tuple[str | None, list[dict[str, Any]]]:
    if not bool(_safe_attr(shape, "has_text_frame", False)):
        return None, []
    frame = shape.text_frame
    return frame.text or None, [_paragraph_data(p) for p in frame.paragraphs]


def _shape_elements(shapes: Iterable[Any], slide_index: int, start_order: int = 1, parent: str | None = None) -> tuple[list[DocumentElement], int, list[UnsupportedObject]]:
    elements: list[DocumentElement] = []
    unsupported: list[UnsupportedObject] = []
    order = start_order
    for z_index, shape in enumerate(shapes, start=1):
        shape_id = _safe_attr(shape, "shape_id")
        element_id = f"slide-{slide_index}-shape-{shape_id or z_index}-{order}"
        source = SourceReference(unit_index=slide_index, slide_number=slide_index)
        shape_type = _enum_name(_safe_attr(shape, "shape_type"))
        base_data = {
            "role": "shape",
            "shape_id": shape_id,
            "name": _safe_attr(shape, "name"),
            "shape_type": shape_type,
            "z_index": z_index,
            "parent_shape": parent,
            "geometry_emu": _shape_geometry(shape),
            "placeholder": _placeholder_data(shape),
            "accessibility": _shape_accessibility(shape),
        }

        if bool(_safe_attr(shape, "has_table", False)):
            table = _table_data(shape.table)
            text = "\n".join(" | ".join(cell["text"] for cell in row) for row in table["rows"])
            elements.append(DocumentElement(element_id, ElementType.TABLE, order, text or None, data={**base_data, "table": table}, source=source))
            order += 1
            continue

        if bool(_safe_attr(shape, "has_chart", False)):
            chart = _chart_data(shape.chart)
            elements.append(DocumentElement(element_id, ElementType.OTHER, order, chart.get("title"), data={**base_data, "role": "chart", "chart": chart}, source=source))
            order += 1
            continue

        image = _safe_attr(shape, "image")
        if image is not None:
            elements.append(DocumentElement(
                element_id, ElementType.IMAGE, order, None,
                data={
                    **base_data,
                    "image": {
                        "filename": _safe_attr(image, "filename"),
                        "extension": _safe_attr(image, "ext"),
                        "content_type": _safe_attr(image, "content_type"),
                        "size_pixels": list(_safe_attr(image, "size", ())) or None,
                        "dpi": list(_safe_attr(image, "dpi", ())) or None,
                        "sha1": _safe_attr(image, "sha1"),
                        "sha256": hashlib.sha256(_safe_attr(image, "blob", b"")).hexdigest()
                        if _safe_attr(image, "blob", None) is not None else None,
                    },
                },
                source=source,
            ))
            order += 1
            continue

        group_shapes = _safe_attr(shape, "shapes")
        if group_shapes is not None and shape_type == "GROUP":
            elements.append(DocumentElement(element_id, ElementType.OTHER, order, None, data={**base_data, "role": "group"}, source=source))
            order += 1
            nested, order, nested_unsupported = _shape_elements(group_shapes, slide_index, order, parent=element_id)
            elements.extend(nested)
            unsupported.extend(nested_unsupported)
            continue

        text, paragraphs = _shape_text(shape)
        if text is not None or paragraphs:
            element_type = ElementType.HEADING if (base_data.get("placeholder") or {}).get("type") in {"TITLE", "CENTER_TITLE", "SUBTITLE"} else ElementType.PARAGRAPH
            elements.append(DocumentElement(element_id, element_type, order, text, data={**base_data, "paragraphs": paragraphs}, source=source))
            order += 1
            continue

        ole_format = _safe_attr(shape, "ole_format")
        if ole_format is not None:
            unsupported.append(UnsupportedObject(
                object_type="pptx_ole_object",
                status="detected_not_fully_parsed",
                description="An embedded OLE object was detected; binary package content is preserved but not interpreted.",
                source=source,
                metadata={**base_data, "prog_id": _safe_attr(ole_format, "prog_id")},
            ))

        elements.append(DocumentElement(element_id, ElementType.OTHER, order, None, data=base_data, source=source))
        order += 1
    return elements, order, unsupported


class PptxExtractor(BaseExtractor):
    name = "python-pptx-native"
    supported_types = frozenset({"pptx", "pptm"})

    def __init__(self) -> None:
        try:
            import pptx
            self.version = pptx.__version__
        except Exception:
            self.version = None

    def extract(self, path: Path, detected: DetectedFileType, metadata: DocumentMetadata) -> CanonicalDocument:
        try:
            presentation = Presentation(path)
        except Exception as exc:
            raise ExtractionFailedError(f"Failed to open PowerPoint presentation: {path.name}: {exc}") from exc

        units: list[DocumentUnit] = []
        warnings: list[ExtractionWarning] = []
        unsupported: list[UnsupportedObject] = []
        masters = []
        for master_index, master in enumerate(presentation.slide_masters, start=1):
            masters.append({
                "index": master_index,
                "name": _safe_attr(master, "name"),
                "visible_text": _shape_text_snapshot(master.shapes),
                "layouts": [
                    {
                        "index": layout_index,
                        "name": _safe_attr(layout, "name"),
                        "visible_text": _shape_text_snapshot(layout.shapes),
                    }
                    for layout_index, layout in enumerate(master.slide_layouts, start=1)
                ],
            })
        metadata.properties["powerpoint"] = {
            "core_properties": _core_properties(presentation.core_properties),
            "slide_count": len(presentation.slides),
            "slide_width_emu": int(presentation.slide_width),
            "slide_height_emu": int(presentation.slide_height),
            "masters": masters,
        }

        for slide_index, slide in enumerate(presentation.slides, start=1):
            elements, order, slide_unsupported = _shape_elements(slide.shapes, slide_index)
            unsupported.extend(slide_unsupported)
            notes_text = None
            if bool(_safe_attr(slide, "has_notes_slide", False)):
                try:
                    frame = slide.notes_slide.notes_text_frame
                    notes_text = frame.text if frame is not None else None
                except Exception as exc:
                    warnings.append(ExtractionWarning(
                        code="pptx.notes.partial",
                        message="Slide notes were detected but could not be fully read.",
                        source=SourceReference(unit_index=slide_index, slide_number=slide_index),
                        details={"error": str(exc)},
                    ))
            if notes_text:
                elements.append(DocumentElement(
                    f"slide-{slide_index}-notes", ElementType.NOTE, order, notes_text,
                    data={"role": "speaker_notes"},
                    source=SourceReference(unit_index=slide_index, slide_number=slide_index),
                ))
                order += 1

            hidden = False
            try:
                show = slide._element.get("show")
                hidden = show in {"0", "false", "False"}
            except Exception:
                pass
            units.append(DocumentUnit(
                index=slide_index,
                kind=UnitKind.SLIDE,
                name=_safe_attr(slide, "name") or f"Slide {slide_index}",
                elements=elements,
                metadata={
                    "slide_id": _safe_attr(slide, "slide_id"),
                    "hidden": hidden,
                    "layout_name": _safe_attr(_safe_attr(slide, "slide_layout"), "name"),
                    "layout_visible_text": _shape_text_snapshot(slide.slide_layout.shapes),
                    "master_visible_text": _shape_text_snapshot(slide.slide_layout.slide_master.shapes),
                    "shape_count": len(slide.shapes),
                    "has_notes": bool(notes_text),
                },
            ))

        archive_names: set[str] = set()
        try:
            with zipfile.ZipFile(path) as archive:
                archive_names = set(archive.namelist())
        except Exception:
            pass
        for prefix, object_type, description in (
            ("ppt/diagrams/", "pptx_smartart", "SmartArt/diagram parts were detected; visual graph structure is not normalized yet."),
            ("ppt/embeddings/", "pptx_embedded_object", "Embedded package/OLE parts were detected; binary content is preserved but not interpreted."),
            ("ppt/comments/", "pptx_comments", "PowerPoint comment parts were detected; comment threading is not normalized yet."),
        ):
            members = sorted(name for name in archive_names if name.startswith(prefix) and not name.endswith("/"))
            if members:
                unsupported.append(UnsupportedObject(
                    object_type=object_type,
                    status="detected_not_fully_parsed",
                    description=description,
                    metadata={"count": len(members), "archive_parts": members},
                ))
        if "ppt/vbaProject.bin" in archive_names:
            unsupported.append(UnsupportedObject(
                object_type="vba_project", status="detected_not_parsed",
                description="A PowerPoint VBA project was detected; macro source is not interpreted.",
                metadata={"archive_part": "ppt/vbaProject.bin"},
            ))

        assets = collect_ooxml_assets(path, detected.document_type)
        asset_by_sha = {asset.sha256: asset.asset_id for asset in assets if asset.sha256}
        for unit in units:
            for element in unit.elements:
                if element.element_type != ElementType.IMAGE:
                    continue
                image = element.data.get("image", {})
                if image.get("sha256") in asset_by_sha:
                    image["asset_id"] = asset_by_sha[image["sha256"]]
        return CanonicalDocument(
            metadata=metadata,
            units=units,
            warnings=warnings,
            unsupported_objects=unsupported,
            assets=assets,
        )

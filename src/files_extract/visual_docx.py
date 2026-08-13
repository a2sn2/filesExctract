from __future__ import annotations

import tempfile
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Pt

from .errors import FilesExtractError

_MAX_WORD_POINTS = 22.0 * 72.0


def pdf_to_visual_docx(source: Path, output: Path, *, dpi: int = 216, password: str | None = None) -> int:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(str(source), password=password)
    except Exception as exc:
        raise FilesExtractError(f"Could not open PDF for conversion: {source.name}: {exc}") from exc

    count = len(pdf)
    if count < 1:
        pdf.close()
        raise FilesExtractError("PDF contains no pages.")

    rendered: list[tuple[Path, float, float]] = []
    with tempfile.TemporaryDirectory(prefix="files-extract-pages-") as tmp:
        root = Path(tmp)
        scale = max(72, min(int(dpi), 600)) / 72.0
        try:
            for index in range(count):
                page = pdf[index]
                width_pt, height_pt = page.get_size()
                render_scale = min(scale, 6000.0 / max(float(width_pt), float(height_pt)))
                bitmap = page.render(scale=max(0.25, render_scale))
                image = bitmap.to_pil()
                page_path = root / f"page-{index + 1}.png"
                image.save(page_path, "PNG")
                image.close()
                rendered.append((page_path, float(width_pt), float(height_pt)))
        finally:
            pdf.close()

        _write_visual_word(rendered, output)
    return count


def image_to_visual_docx(source: Path, output: Path) -> int:
    with Image.open(source) as image:
        image.load()
        dpi = image.info.get("dpi", (96.0, 96.0))
        try:
            dpi_x, dpi_y = float(dpi[0]), float(dpi[1])
            if dpi_x <= 1 or dpi_y <= 1:
                raise ValueError
        except Exception:
            dpi_x = dpi_y = 96.0
        width_pt = image.width * 72.0 / dpi_x
        height_pt = image.height * 72.0 / dpi_y

    _write_visual_word([(source, width_pt, height_pt)], output)
    return 1


def _write_visual_word(pages: list[tuple[Path, float, float]], output: Path) -> None:
    if not pages:
        raise FilesExtractError("No pages were available for Word conversion.")

    document = Document()
    for index, (image_path, width_pt, height_pt) in enumerate(pages):
        size_scale = min(1.0, _MAX_WORD_POINTS / max(width_pt, height_pt))
        width = width_pt * size_scale
        height = height_pt * size_scale

        if index == 0:
            section = document.sections[0]
            paragraph = document.paragraphs[0]
        else:
            section = document.add_section(WD_SECTION.NEW_PAGE)
            paragraph = document.add_paragraph()

        section.page_width = Pt(width)
        section.page_height = Pt(height)
        section.top_margin = Pt(0)
        section.bottom_margin = Pt(0)
        section.left_margin = Pt(0)
        section.right_margin = Pt(0)
        section.header_distance = Pt(0)
        section.footer_distance = Pt(0)

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run().add_picture(str(image_path), width=Pt(width), height=Pt(height))

    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "FilesExtract Exact Layout Conversion"
    document.save(output)
